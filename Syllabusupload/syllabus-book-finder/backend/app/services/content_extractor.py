import os
import re
from pathlib import Path
from typing import List, Dict, Any, Optional
import docx
import pptx
import fitz  # PyMuPDF

class DocumentChunk:
    def __init__(
        self,
        chunk_id: str,
        slide_or_page_number: int,
        title: str,
        content: str,
        headings: Optional[List[str]] = None,
        bullet_points: Optional[List[str]] = None,
        code_snippets: Optional[List[str]] = None,
        tables: Optional[List[str]] = None,
        is_image_only: bool = False,
        character_count: int = 0
    ):
        self.chunk_id = chunk_id
        self.slide_or_page_number = slide_or_page_number
        self.title = title
        self.content = content
        self.headings = headings or []
        self.bullet_points = bullet_points or []
        self.code_snippets = code_snippets or []
        self.tables = tables or []
        self.is_image_only = is_image_only
        self.character_count = character_count or len(content)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "chunk_id": self.chunk_id,
            "slide_or_page_number": self.slide_or_page_number,
            "title": self.title,
            "content": self.content,
            "headings": self.headings,
            "bullet_points": self.bullet_points,
            "code_snippets": self.code_snippets,
            "tables": self.tables,
            "is_image_only": self.is_image_only,
            "character_count": self.character_count
        }

class ExtractionResult:
    def __init__(
        self,
        text: str,
        slide_or_page_count: int,
        chunks: List[Dict[str, Any]],
        status: str = "extracted",
        warnings: Optional[List[str]] = None
    ):
        self.text = text
        self.slide_or_page_count = slide_or_page_count
        self.chunks = chunks
        self.status = status
        self.warnings = warnings or []

def filter_repeated_headers_footers(chunks: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Detect and remove identical headers and footers appearing across multiple pages/slides."""
    if len(chunks) < 3:
        return chunks

    line_counts = {}
    for c in chunks:
        lines = [l.strip() for l in c["content"].splitlines() if l.strip()]
        if lines:
            first_line = lines[0]
            last_line = lines[-1]
            line_counts[first_line] = line_counts.get(first_line, 0) + 1
            if len(lines) > 1:
                line_counts[last_line] = line_counts.get(last_line, 0) + 1

    threshold = max(2, int(len(chunks) * 0.45))
    repeated_lines = {line for line, count in line_counts.items() if count >= threshold and len(line) < 80}

    for c in chunks:
        lines = c["content"].splitlines()
        filtered = [l for l in lines if l.strip() not in repeated_lines]
        c["content"] = "\n".join(filtered).strip()
        c["character_count"] = len(c["content"])

    return chunks

def extract_from_pptx(file_path: Path) -> ExtractionResult:
    prs = pptx.Presentation(str(file_path))
    slide_count = len(prs.slides)
    chunks = []
    warnings = []
    all_text_blocks = []

    for idx, slide in enumerate(prs.slides, start=1):
        slide_title = ""
        headings = []
        bullets = []
        code_snippets = []
        table_rows = []
        slide_lines = []
        has_images = False

        if slide.shapes.title and slide.shapes.title.text:
            slide_title = slide.shapes.title.text.strip()
            headings.append(slide_title)
            slide_lines.append(f"Title: {slide_title}")

        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue

            # Detect picture/image shapes
            try:
                if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
                    has_images = True
            except Exception:
                pass

            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = p.text.strip()
                    if not t or t == slide_title:
                        continue
                    
                    # Detect code snippets or algorithms
                    if re.search(r'[{};=><\+\-\*/]|\b(function|def|return|for|while|if|else|import|library|ggplot)\b', t):
                        code_snippets.append(t)
                    elif t.startswith(("-", "•", "*", "1.", "2.", "3.", "a.", "b.")):
                        bullets.append(t)
                    else:
                        slide_lines.append(t)

            elif shape.has_table:
                for row in shape.table.rows:
                    row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if row_cells:
                        table_rows.append(" | ".join(row_cells))

        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                slide_lines.append(f"Speaker Notes: {notes}")

        combined_content = "\n".join(headings + slide_lines + bullets + code_snippets + table_rows).strip()
        is_image_only = (len(combined_content) < 25 and has_images)

        if is_image_only:
            warnings.append(f"Slide {idx}: Contains image/diagram with minimal text. OCR may be required for full extraction.")

        chunk = DocumentChunk(
            chunk_id=f"slide_{idx}",
            slide_or_page_number=idx,
            title=slide_title or f"Slide {idx}",
            content=combined_content or f"[Slide {idx} - Visual Diagram / Minimal Text]",
            headings=headings,
            bullet_points=bullets,
            code_snippets=code_snippets,
            tables=table_rows,
            is_image_only=is_image_only,
            character_count=len(combined_content)
        )
        chunks.append(chunk.to_dict())
        if combined_content:
            all_text_blocks.append(f"--- Slide {idx}: {slide_title or 'Content'} ---\n{combined_content}")

    chunks = filter_repeated_headers_footers(chunks)
    total_text = "\n\n".join(all_text_blocks).strip()
    status = "partially_extracted" if warnings and len(warnings) > slide_count // 2 else ("failed" if not total_text else "extracted")

    return ExtractionResult(
        text=total_text,
        slide_or_page_count=slide_count,
        chunks=chunks,
        status=status,
        warnings=warnings
    )

def extract_from_pdf(file_path: Path) -> ExtractionResult:
    doc = fitz.open(str(file_path))
    page_count = len(doc)
    chunks = []
    warnings = []
    all_text_blocks = []

    for page_num in range(page_count):
        idx = page_num + 1
        page = doc[page_num]
        raw_text = page.get_text("text").strip()
        image_list = page.get_images()
        has_images = len(image_list) > 0

        headings = []
        bullets = []
        code_snippets = []
        tables = []
        page_lines = []

        lines = [l.strip() for l in raw_text.splitlines() if l.strip()]
        page_title = lines[0] if lines else f"Page {idx}"

        for line in lines:
            if re.match(r'^(?:Chapter|Unit|Module|Section|[0-9]+[\.\)])\s+', line, re.IGNORECASE) or (line.isupper() and len(line) < 60):
                headings.append(line)
            elif line.startswith(("-", "•", "*", "–")):
                bullets.append(line)
            elif re.search(r'[{};=><\+\-\*/]|\b(def|function|return|import|class|library)\b', line):
                code_snippets.append(line)
            elif "|" in line or "\t" in line:
                tables.append(line)
            else:
                page_lines.append(line)

        is_image_only = (len(raw_text) < 35 and has_images)
        if is_image_only:
            warnings.append(f"Page {idx}: Minimal text detected. Document appears to be a scanned page or diagram; OCR may be required.")

        chunk = DocumentChunk(
            chunk_id=f"page_{idx}",
            slide_or_page_number=idx,
            title=page_title[:80],
            content=raw_text or f"[Page {idx} - Scanned / Image Only]",
            headings=headings,
            bullet_points=bullets,
            code_snippets=code_snippets,
            tables=tables,
            is_image_only=is_image_only,
            character_count=len(raw_text)
        )
        chunks.append(chunk.to_dict())
        if raw_text:
            all_text_blocks.append(f"--- Page {idx} ---\n{raw_text}")

    doc.close()
    chunks = filter_repeated_headers_footers(chunks)
    total_text = "\n\n".join(all_text_blocks).strip()
    status = "partially_extracted" if warnings and len(warnings) > page_count // 2 else ("failed" if not total_text else "extracted")

    return ExtractionResult(
        text=total_text,
        slide_or_page_count=page_count,
        chunks=chunks,
        status=status,
        warnings=warnings
    )

def extract_from_docx(file_path: Path) -> ExtractionResult:
    doc = docx.Document(str(file_path))
    chunks = []
    warnings = []
    all_text_blocks = []

    current_title = "Document Introduction"
    current_lines = []
    current_headings = []
    current_bullets = []
    current_code = []
    current_tables = []
    page_counter = 1

    def flush_chunk():
        nonlocal page_counter, current_title, current_lines, current_headings, current_bullets, current_code, current_tables
        content = "\n".join(current_headings + current_lines + current_bullets + current_code + current_tables).strip()
        if content:
            chunks.append(DocumentChunk(
                chunk_id=f"section_{page_counter}",
                slide_or_page_number=page_counter,
                title=current_title[:80],
                content=content,
                headings=list(current_headings),
                bullet_points=list(current_bullets),
                code_snippets=list(current_code),
                tables=list(current_tables),
                is_image_only=False,
                character_count=len(content)
            ).to_dict())
            all_text_blocks.append(f"--- Section / Page {page_counter}: {current_title} ---\n{content}")
            page_counter += 1
            current_lines = []
            current_headings = []
            current_bullets = []
            current_code = []
            current_tables = []

    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue

        if p.style.name.startswith("Heading") or (len(t) < 60 and (t.isupper() or re.match(r'^[0-9]+[\.\)]\s+[A-Z]', t))):
            if len(current_lines) > 4 or len(current_bullets) > 3:
                flush_chunk()
            current_title = t
            current_headings.append(t)
        elif t.startswith(("-", "•", "*", "–")) or p.style.name.startswith("List"):
            current_bullets.append(t)
        elif re.search(r'[{};=><\+\-\*/]|\b(def|function|return|import|class|library)\b', t):
            current_code.append(t)
        else:
            current_lines.append(t)

    for table in doc.tables:
        t_rows = []
        for row in table.rows:
            r_cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if r_cells:
                t_rows.append(" | ".join(r_cells))
        if t_rows:
            current_tables.extend(t_rows)
            if len(current_tables) > 10:
                flush_chunk()

    flush_chunk()
    total_text = "\n\n".join(all_text_blocks).strip()
    return ExtractionResult(
        text=total_text,
        slide_or_page_count=max(1, len(chunks)),
        chunks=chunks,
        status="extracted" if total_text else "failed",
        warnings=warnings
    )

def extract_from_txt(file_path: Path) -> ExtractionResult:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        content = f.read()

    lines = [l.strip() for l in content.splitlines() if l.strip()]
    chunks = []
    chunk_size = 35
    for idx, i in enumerate(range(0, len(lines), chunk_size), start=1):
        batch = lines[i:i+chunk_size]
        title = batch[0][:60] if batch else f"Section {idx}"
        chunk_text = "\n".join(batch)
        chunks.append(DocumentChunk(
            chunk_id=f"part_{idx}",
            slide_or_page_number=idx,
            title=title,
            content=chunk_text,
            character_count=len(chunk_text)
        ).to_dict())

    return ExtractionResult(
        text=content,
        slide_or_page_count=max(1, len(chunks)),
        chunks=chunks,
        status="extracted" if content.strip() else "failed"
    )

def extract_document_content(file_path: Path) -> ExtractionResult:
    """Universal document extractor with slide/page preservation and OCR detection."""
    ext = file_path.suffix.lower()
    if ext in [".pptx", ".ppt"]:
        if ext == ".pptx":
            return extract_from_pptx(file_path)
        else:
            from app.services.pptx_extractor import _extract_from_legacy_ppt
            res = _extract_from_legacy_ppt(file_path)
            return ExtractionResult(
                text=res.text,
                slide_or_page_count=res.slide_count,
                chunks=[{"chunk_id": "slide_1", "slide_or_page_number": 1, "title": "Legacy PPT", "content": res.text, "character_count": len(res.text)}],
                status="extracted" if res.text else "failed"
            )
    elif ext == ".pdf":
        return extract_from_pdf(file_path)
    elif ext == ".docx":
        return extract_from_docx(file_path)
    elif ext in [".txt", ".md"]:
        return extract_from_txt(file_path)
    else:
        raise ValueError(f"Unsupported file format '{ext}'. Allowed: .pptx, .ppt, .pdf, .docx, .txt")
