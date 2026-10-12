import pytest
from pathlib import Path
import docx
from docx.shared import Pt, Inches
import pandas as pd
from app.services.template_extractor import (
    extract_blueprint_from_docx,
    extract_blueprint_from_xlsx,
    validate_extracted_blueprint,
    normalize_bloom,
    normalize_co,
    parse_duration_to_minutes
)

@pytest.fixture
def sample_docx_template(tmp_path: Path) -> Path:
    doc_path = tmp_path / "test_template.docx"
    doc = docx.Document()
    
    # Institution Header
    p = doc.add_paragraph("AARUPADAI VEEDU INSTITUTE OF TECHNOLOGY")
    p.runs[0].bold = True
    doc.add_paragraph("DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING")
    doc.add_paragraph("Internal Assessment - 2024-2025")
    doc.add_paragraph("Course Code: 35021P28 | Course Title: R Programming")
    doc.add_paragraph("Max. Marks: 50 | Duration: 90 Minutes")
    doc.add_paragraph("Instructions: 1. Answer all questions in Part A. 2. Answer either (a) or (b) in Part B.")

    # Part A Table
    doc.add_paragraph("PART - A")
    tbl_a = doc.add_table(rows=1, cols=5)
    hdr_a = tbl_a.rows[0].cells
    hdr_a[0].text = "Q.No"
    hdr_a[1].text = "Questions"
    hdr_a[2].text = "CO"
    hdr_a[3].text = "Bloom's Level"
    hdr_a[4].text = "Marks"

    for i in range(1, 6):
        row = tbl_a.add_row().cells
        row[0].text = str(i)
        row[1].text = f"Explain concept {i} in R programming."
        row[2].text = f"CO{(i % 2) + 1}"
        row[3].text = "L2"
        row[4].text = "2"

    # Part B Table (with OR choices)
    doc.add_paragraph("PART - B")
    tbl_b = doc.add_table(rows=1, cols=5)
    hdr_b = tbl_b.rows[0].cells
    hdr_b[0].text = "Q.No"
    hdr_b[1].text = "Questions"
    hdr_b[2].text = "CO"
    hdr_b[3].text = "Bloom's Level"
    hdr_b[4].text = "Marks"

    # Q6(a)
    r1 = tbl_b.add_row().cells
    r1[0].text = "6(a)"
    r1[1].text = "Write an R program to compute vector mean."
    r1[2].text = "CO2"
    r1[3].text = "L3"
    r1[4].text = "10"

    # Divider OR
    r_or = tbl_b.add_row().cells
    r_or[0].text = ""
    r_or[1].text = "[OR]"
    r_or[2].text = ""
    r_or[3].text = ""
    r_or[4].text = ""

    # Q6(b)
    r2 = tbl_b.add_row().cells
    r2[0].text = "6(b)"
    r2[1].text = "Write an R script to sort matrices by row."
    r2[2].text = "CO2"
    r2[3].text = "L3"
    r2[4].text = "10"

    # Q7(a)
    r3 = tbl_b.add_row().cells
    r3[0].text = "7(a)"
    r3[1].text = "Describe factor variables and data frames in R."
    r3[2].text = "CO3"
    r3[3].text = "L3"
    r3[4].text = "10"

    # Divider OR
    r_or2 = tbl_b.add_row().cells
    r_or2[0].text = ""
    r_or2[1].text = "[OR]"
    r_or2[2].text = ""
    r_or2[3].text = ""
    r_or2[4].text = ""

    # Q7(b)
    r4 = tbl_b.add_row().cells
    r4[0].text = "7(b)"
    r4[1].text = "Explain the apply family of functions in R."
    r4[2].text = "CO3"
    r4[3].text = "L3"
    r4[4].text = "10"

    # Q8 (Part C / Long question)
    doc.add_paragraph("PART - C")
    tbl_c = doc.add_table(rows=1, cols=5)
    hdr_c = tbl_c.rows[0].cells
    hdr_c[0].text = "Q.No"
    hdr_c[1].text = "Questions"
    hdr_c[2].text = "CO"
    hdr_c[3].text = "Bloom's Level"
    hdr_c[4].text = "Marks"

    r_c1 = tbl_c.add_row().cells
    r_c1[0].text = "8"
    r_c1[1].text = "Design a machine learning classification pipeline in R."
    r_c1[2].text = "CO4"
    r_c1[3].text = "L5"
    r_c1[4].text = "20"

    doc.save(str(doc_path))
    return doc_path

@pytest.fixture
def sample_xlsx_template(tmp_path: Path) -> Path:
    xlsx_path = tmp_path / "test_blueprint.xlsx"
    data = [
        {"Section": "Part A", "Q.No": "1", "Unit": 1, "CO": "CO1", "Bloom": "L2", "Marks": 2, "Type": "mcq_single"},
        {"Section": "Part A", "Q.No": "2", "Unit": 1, "CO": "CO1", "Bloom": "L2", "Marks": 2, "Type": "mcq_single"},
        {"Section": "Part A", "Q.No": "3", "Unit": 2, "CO": "CO2", "Bloom": "L2", "Marks": 2, "Type": "mcq_single"},
        {"Section": "Part A", "Q.No": "4", "Unit": 2, "CO": "CO2", "Bloom": "L2", "Marks": 2, "Type": "mcq_single"},
        {"Section": "Part A", "Q.No": "5", "Unit": 3, "CO": "CO3", "Bloom": "L2", "Marks": 2, "Type": "mcq_single"},
        {"Section": "Part B", "Q.No": "6(a)", "Unit": 2, "CO": "CO2", "Bloom": "L3", "Marks": 10, "Type": "coding,descriptive"},
        {"Section": "Part B", "Q.No": "6(b)", "Unit": 2, "CO": "CO2", "Bloom": "L3", "Marks": 10, "Type": "coding,descriptive"},
        {"Section": "Part B", "Q.No": "7(a)", "Unit": 3, "CO": "CO3", "Bloom": "L3", "Marks": 10, "Type": "coding,descriptive"},
        {"Section": "Part B", "Q.No": "7(b)", "Unit": 3, "CO": "CO3", "Bloom": "L3", "Marks": 10, "Type": "coding,descriptive"},
        {"Section": "Part C", "Q.No": "8", "Unit": 4, "CO": "CO4", "Bloom": "L5", "Marks": 20, "Type": "case_study,design"}
    ]
    df = pd.DataFrame(data)
    df.to_excel(str(xlsx_path), index=False)
    return xlsx_path

def test_docx_template_extraction(sample_docx_template: Path):
    blueprint, conf, full_text = extract_blueprint_from_docx(sample_docx_template)

    assert "AARUPADAI VEEDU" in blueprint["institution"]
    assert blueprint["course_code"] == "35021P28"
    assert blueprint["maximum_marks"] == 50
    assert blueprint["duration_minutes"] == 90
    assert len(blueprint["sections"]) >= 2

    # Part A verification
    sec_a = blueprint["sections"][0]
    assert sec_a["marks_per_question"] == 2
    assert len(sec_a["questions"]) == 5

    # Part B verification with OR choices
    sec_b = blueprint["sections"][1]
    assert sec_b["or_choice"] is True
    assert any(q.get("or_pair_id") for q in sec_b["questions"])

    # Confidence check
    assert conf["institution"] >= 0.8
    assert conf["overall"] >= 0.8

def test_xlsx_template_extraction(sample_xlsx_template: Path):
    blueprint, conf, text = extract_blueprint_from_xlsx(sample_xlsx_template)

    assert blueprint["maximum_marks"] == 50
    assert len(blueprint["sections"]) == 3
    assert conf["overall"] >= 0.9

    sec_b = blueprint["sections"][1]
    assert sec_b["or_choice"] is True
    assert sec_b["marks_per_question"] == 10

def test_blueprint_validation_rules():
    # Valid 50-mark blueprint
    valid_bp = {
        "institution": "AVIT",
        "course_code": "35021P28",
        "course_title": "R Programming",
        "maximum_marks": 50,
        "sections": [
            {
                "section_name": "Part A",
                "answer_rule": "Answer all questions",
                "section_total": 10,
                "or_choice": False,
                "questions": [
                    {"question_number": "1", "marks": 2, "bloom_level": "L2", "co": "CO1", "question_types": ["mcq_single"]},
                    {"question_number": "2", "marks": 2, "bloom_level": "L2", "co": "CO1", "question_types": ["mcq_single"]},
                    {"question_number": "3", "marks": 2, "bloom_level": "L2", "co": "CO2", "question_types": ["mcq_single"]},
                    {"question_number": "4", "marks": 2, "bloom_level": "L2", "co": "CO2", "question_types": ["mcq_single"]},
                    {"question_number": "5", "marks": 2, "bloom_level": "L2", "co": "CO3", "question_types": ["mcq_single"]}
                ]
            },
            {
                "section_name": "Part B",
                "answer_rule": "Answer either (a) or (b)",
                "section_total": 20,
                "or_choice": True,
                "questions": [
                    {"question_number": "6(a)", "or_pair_id": "or_pair_6", "marks": 10, "bloom_level": "L3", "co": "CO2", "question_types": ["coding"]},
                    {"question_number": "6(b)", "or_pair_id": "or_pair_6", "marks": 10, "bloom_level": "L3", "co": "CO2", "question_types": ["coding"]},
                    {"question_number": "7(a)", "or_pair_id": "or_pair_7", "marks": 10, "bloom_level": "L3", "co": "CO3", "question_types": ["coding"]},
                    {"question_number": "7(b)", "or_pair_id": "or_pair_7", "marks": 10, "bloom_level": "L3", "co": "CO3", "question_types": ["coding"]}
                ]
            },
            {
                "section_name": "Part C",
                "answer_rule": "Answer all questions",
                "section_total": 20,
                "or_choice": False,
                "questions": [
                    {"question_number": "8", "marks": 20, "bloom_level": "L5", "co": "CO4", "question_types": ["case_study"]}
                ]
            }
        ]
    }

    res = validate_extracted_blueprint(valid_bp)
    assert res["is_valid"] is True
    assert res["reconciled"] is True
    assert len(res["errors"]) == 0

    # Test Parity of Marks Failure in OR pair
    invalid_or_bp = dict(valid_bp)
    invalid_or_bp["sections"][1]["questions"][1]["marks"] = 8 # Unequal marks: 10 vs 8
    res_inv = validate_extracted_blueprint(invalid_or_bp)
    assert res_inv["is_valid"] is False
    assert any("equal marks" in e for e in res_inv["errors"])

    # Test Total Marks Mismatch
    mismatch_bp = dict(valid_bp)
    mismatch_bp["maximum_marks"] = 100 # Target 100 but sections sum to 50
    res_mis = validate_extracted_blueprint(mismatch_bp)
    assert res_mis["is_valid"] is False
    assert any("reconcile" in e for e in res_mis["errors"])

def test_normalization_helpers():
    assert normalize_bloom("K3 - Apply") == "L3"
    assert normalize_bloom("L4") == "L4"
    assert normalize_bloom("understand") == "L2"
    assert normalize_co("Course Outcome 3") == "CO3"
    assert normalize_co("CO5") == "CO5"
    assert parse_duration_to_minutes("3 Hours") == 180
    assert parse_duration_to_minutes("90 Minutes") == 90
    assert parse_duration_to_minutes("1.5 hr") == 90
