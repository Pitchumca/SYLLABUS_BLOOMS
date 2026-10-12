import io
import pytest
import openpyxl
from app.services.export_service import (
    generate_excel_report,
    generate_pdf_report,
    generate_csv_report,
    generate_json_report
)

@pytest.fixture
def dummy_books():
    return [
        {
            "category": "Textbook",
            "title": "Introduction to Algorithms",
            "authors": "Thomas H. Cormen et al.",
            "edition": "3rd Edition",
            "publisher": "MIT Press",
            "publication_year": "2009",
            "isbn": "9780262033848",
            "source": "Google Books",
            "access_type": "Preview available",
            "access_url": "https://books.google.com/sample",
            "confidence_score": 95.0,
            "verification_status": "Verified API match",
            "remarks": "Sample chapters available."
        },
        {
            "category": "Reference Book",
            "title": "Data Structures in Python",
            "authors": "Michael T. Goodrich",
            "edition": "1st Edition",
            "publisher": "Wiley",
            "publication_year": "2013",
            "isbn": "9781118290279",
            "source": "Open Library",
            "access_type": "Borrowing available",
            "access_url": "https://archive.org/details/sample",
            "confidence_score": 88.0,
            "verification_status": "Controlled digital lending",
            "remarks": "14-day borrowing available."
        }
    ]

def test_generate_excel_report(dummy_books):
    excel_bytes = generate_excel_report("CS401 Syllabus", dummy_books)
    assert len(excel_bytes) > 1000
    
    # Verify openpyxl can load it
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Book Access Audit" in wb.sheetnames
    ws = wb["Book Access Audit"]
    assert "AI-Powered Syllabus" in str(ws["A1"].value)
    assert ws["A7"].value == "S.No."
    assert ws["B8"].value == "Textbook"
    assert ws["C8"].value == "Introduction to Algorithms"

def test_generate_pdf_report(dummy_books):
    pdf_bytes = generate_pdf_report("CS401 Syllabus", dummy_books)
    assert len(pdf_bytes) > 2000
    assert pdf_bytes.startswith(b"%PDF")

def test_generate_csv_report(dummy_books):
    csv_str = generate_csv_report(dummy_books)
    assert "S.No.,Category,Book Title" in csv_str
    assert "Introduction to Algorithms" in csv_str
    assert "Data Structures in Python" in csv_str

def test_generate_json_report(dummy_books):
    import json
    json_str = generate_json_report("CS401 Syllabus", dummy_books)
    parsed = json.loads(json_str)
    assert parsed["metadata"]["syllabus_name"] == "CS401 Syllabus"
    assert len(parsed["books"]) == 2
