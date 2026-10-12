import pytest
import json
import io
import docx
from app.services.grammar_validator import (
    clean_academic_language,
    validate_question_grammar,
    BANNED_ENTERPRISE_PHRASES,
    SYLLABUS_PHRASE_CORRECTIONS
)
from app.services.semantic_validator import (
    evaluate_semantic_quality,
    evaluate_or_pair_equivalence,
    validate_mcq_quality
)
from app.services.question_generator import (
    generate_questions_from_blueprint,
    R_PROGRAMMING_QUESTION_BANK
)
from app.services.word_template_service import generate_question_paper_docx


def test_grammar_validator_fixes_awkward_phrasing():
    """Verify that awkward syllabus phrases like 'Introducing to R' and 'accessing keyboard and monitor' are normalized."""
    text_with_introducing = "Explain the fundamentals of Introducing to R with simple examples."
    cleaned1 = clean_academic_language(text_with_introducing)
    assert "Introducing to R" not in cleaned1
    assert "Introduction to R" in cleaned1

    text_with_keyboard = "Describe the mechanism of accessing keyboard and monitor in R scripts."
    cleaned2 = clean_academic_language(text_with_keyboard)
    assert "accessing keyboard and monitor" not in cleaned2
    assert "reading input from the keyboard and displaying output on the monitor" in cleaned2


def test_grammar_validator_purges_enterprise_jargon():
    """Verify that generic IT buzzwords are completely stripped and penalized."""
    jargon_question = (
        "Evaluate the architectural viability of Generalized Linear Models for mission-critical enterprise environments "
        "to fulfil enterprise scalability and fault-tolerance criteria."
    )
    cleaned, score, issues = validate_question_grammar(jargon_question)
    for phrase in [
        "architectural viability",
        "mission-critical enterprise environments",
        "enterprise scalability",
        "fault-tolerance criteria"
    ]:
        assert phrase not in cleaned.lower()
    
    # Check that jargon penalties triggered issue detections
    assert len(issues) > 0
    assert any("banned enterprise" in iss.lower() for iss in issues)


def test_mcq_option_validation_and_rejection():
    """Verify MCQ validation catches bad options (duplicates, missing correct option, etc.)."""
    # Valid MCQ
    valid_mcq = {
        "options_json": json.dumps([
            "A) vector()",
            "B) list()",
            "C) data.frame()",
            "D) matrix()"
        ]),
        "correct_answer_json": "B) list()",
        "explanation": "list() creates generic vectors holding elements of different types."
    }
    score, issues = validate_mcq_quality(valid_mcq)
    assert score >= 90
    assert len(issues) == 0

    # Duplicate options should trigger deductions
    dup_mcq = {
        "options_json": json.dumps([
            "A) vector()",
            "B) vector()",
            "C) data.frame()",
            "D) matrix()"
        ]),
        "correct_answer_json": "B) vector()"
    }
    score_dup, issues_dup = validate_mcq_quality(dup_mcq)
    assert score_dup < 80
    assert any("duplicate" in iss.lower() for iss in issues_dup)

    # Missing correct option should trigger deductions
    missing_key_mcq = {
        "options_json": json.dumps([
            "A) vector()",
            "B) list()",
            "C) data.frame()",
            "D) matrix()"
        ]),
        "correct_answer_json": ""
    }
    score_mk, issues_mk = validate_mcq_quality(missing_key_mcq)
    assert score_mk < 80
    assert any("missing correct answer" in iss.lower() for iss in issues_mk)


def test_semantic_quality_validator_thresholds():
    """Verify semantic validator applies strict thresholds and flags needs_review."""
    good_question = {
        "question_text": "Apply glm() to fit a logistic regression model on binary response data frame in R.",
        "co_level": "CO5",
        "bloom_level": "L3 - Apply",
        "marks": 5,
        "unit_number": 5,
        "question_type": "programming",
        "source_filename": "Unit_V_Models.pptx",
        "supporting_passage": "Logistic regression in R is fitted using glm(formula, family = binomial, data).",
        "source_topic": "Linear and Generalized Linear Models",
        "model_answer": "Use glm(y ~ x1 + x2, data = df, family = binomial(link = 'logit')) to fit logistic regression.",
        "marking_scheme": "2 marks for glm formula, 2 marks for family = binomial, 1 mark for syntax execution."
    }
    unit_info = {"number": "5", "title": "Advanced Analytics and Models", "topics": "Linear and Generalized Linear Models, Clustering"}
    co_info = {"code": "CO5", "description": "Formulate predictive models and evaluate linear/generalized models"}

    res = evaluate_semantic_quality(good_question, unit_info, co_info)
    assert res["status"] == "draft"
    assert res["quality_score"] >= 85
    assert res["grammar_score"] >= 90

    # Off-topic question should fail syllabus/CO check and be marked needs_review
    bad_question = {
        "question_text": "Explain the history of the Java Virtual Machine and bytecodes.",
        "co_level": "CO5",
        "bloom_level": "L3 - Apply",
        "marks": 5,
        "unit_number": 1,  # Conflict: unit 1 vs CO5
        "question_type": "descriptive",
        "source_filename": "",
        "supporting_passage": "",
        "source_topic": "Java Architecture",
        "model_answer": "",
        "marking_scheme": ""
    }
    res_bad = evaluate_semantic_quality(bad_question, unit_info, co_info)
    assert res_bad["status"] == "needs_review"
    assert len(res_bad["reasons"]) > 0


def test_or_pair_equivalence_evaluation():
    """Verify OR-pair equivalence scoring detects matched vs mismatched pairs."""
    q_a = {
        "question_number": "11(a)",
        "marks": 8,
        "unit_number": 2,
        "co_level": "CO2",
        "bloom_level": "L3 - Apply",
        "difficulty": "Medium",
        "question_type": "programming",
        "model_answer": "Step 1: Use filter() to select active accounts. Step 2: mutate() to normalize values."
    }
    q_b_matched = {
        "question_number": "11(b)",
        "marks": 8,
        "unit_number": 2,
        "co_level": "CO2",
        "bloom_level": "L3 - Apply",
        "difficulty": "Medium",
        "question_type": "programming",
        "model_answer": "Step 1: Use merge() to join datasets. Step 2: replace NA values with medians."
    }
    score, summary = evaluate_or_pair_equivalence(q_a, q_b_matched)
    assert score >= 85
    assert "Equivalence: 100" in summary or score >= 85

    # Mismatched unit and bloom level
    q_b_mismatched = {
        "question_number": "11(b)",
        "marks": 8,
        "unit_number": 5,  # Mismatched unit
        "co_level": "CO5",  # Mismatched CO
        "bloom_level": "L6 - Create",  # Mismatched Bloom
        "difficulty": "Hard",
        "question_type": "descriptive",
        "model_answer": "Detailed architecture design."
    }
    score_bad, summary_bad = evaluate_or_pair_equivalence(q_a, q_b_mismatched)
    assert score_bad < 70
    assert "mismatch" in summary_bad.lower() or "divergence" in summary_bad.lower()


def test_part_c_strictly_grounded_in_unit_5():
    """Verify that Part C questions strictly draw from Unit 5 and CO5, never Unit 1."""
    units = [
        {"number": "1", "title": "Introduction to R", "topics": "History • Vectors • Arrays • Matrices"},
        {"number": "2", "title": "Control Structures and Functions", "topics": "if-else • loops • functions • apply"},
        {"number": "3", "title": "Data Manipulation", "topics": "Data frames • Dplyr • Reshape"},
        {"number": "4", "title": "Graphics and Visualization", "topics": "Base graphics • ggplot2 • Lattice"},
        {"number": "5", "title": "Statistical Modeling and Interfacing", "topics": "Linear Models • Generalized Linear Models • Clustering • C++ Interfacing"}
    ]
    cos = [
        {"code": "CO1", "description": "Understand R syntax and basic data structures", "bloom_level": "L2 - Understand"},
        {"code": "CO2", "description": "Write control structures and functions", "bloom_level": "L3 - Apply"},
        {"code": "CO3", "description": "Manipulate tabular datasets", "bloom_level": "L3 - Apply"},
        {"code": "CO4", "description": "Create data visualisations", "bloom_level": "L4 - Analyze"},
        {"code": "CO5", "description": "Build statistical and generalized linear models", "bloom_level": "L6 - Create"}
    ]
    blueprint_sections = [
        {
            "name": "Part C",
            "question_type": "design_application",
            "total_questions": 1,
            "questions_to_answer": 1,
            "marks_per_question": 15,
            "has_internal_choice": False,
            "bloom_distribution": ["L6 - Create"],
            "co_distribution": ["CO5"]
        }
    ]

    blueprint = {"sections": blueprint_sections}
    generated = generate_questions_from_blueprint(
        blueprint=blueprint,
        syllabus_text="",
        course_outcomes=cos,
        units=units,
        aligned_chunks=[]
    )
    assert len(generated) >= 1
    part_c_q = generated[0]
    
    # Must be strictly Unit 5 and CO5
    assert str(part_c_q["unit_number"]) == "5"
    assert part_c_q["co_level"] == "CO5"
    assert "L6" in part_c_q["bloom_level"]
    # Question text should be related to Unit 5 topics (modeling, clustering, interfacing, glm)
    text_lower = part_c_q["question_text"].lower()
    assert any(term in text_lower for term in ["model", "clustering", "glm", "linear", "k-means", "rcpp", "interfac", "predict"])
    # Must NOT have Unit 1 intro topics
    assert "introducing to r" not in text_lower
    assert "history of r" not in text_lower


def test_mcq_word_export():
    """Verify that Word (.docx) export formats MCQs with indented options in exam mode, and includes rationales in answer key mode."""
    sample_paper = {
        "institution_name": "Department of Computer Science",
        "course_name": "R Programming for Data Analytics",
        "course_code": "CS302",
        "programme": "B.Tech Data Science",
        "regulation": "R2024",
        "exam_name": "End Semester Examination",
        "academic_year": "2025-2026",
        "semester": "V",
        "department": "Computer Science & Engineering",
        "duration": "3 Hours",
        "max_marks": 100,
        "instructions": "Answer all questions.",
        "generation_mode": "syllabus_ppt"
    }

    mcq_options = [
        {"label": "A", "text": "vector()", "is_correct": False, "explanation": "vector() produces atomic vectors, not nested elements"},
        {"label": "B", "text": "list()", "is_correct": True, "explanation": "list() creates generic vectors holding elements of different types"},
        {"label": "C", "text": "matrix()", "is_correct": False, "explanation": "matrix() requires all elements to share identical atomic types"},
        {"label": "D", "text": "array()", "is_correct": False, "explanation": "array() is a multi-dimensional homogeneous structure"}
    ]
    distractor_dict = {
        "A": "vector() produces atomic vectors, not nested elements",
        "C": "matrix() requires all elements to share identical atomic types",
        "D": "array() is a multi-dimensional homogeneous structure"
    }

    questions = [
        {
            "id": 1,
            "section_name": "Part A",
            "question_number": "1",
            "question_text": "Which R function creates a heterogeneous container capable of holding elements of varying data types?",
            "question_type": "mcq",
            "marks": 2,
            "unit_number": 1,
            "co_level": "CO1",
            "bloom_level": "L2 - Understand",
            "difficulty": "Easy",
            "options_json": json.dumps(mcq_options),
            "correct_option": "B",
            "distractor_explanations_json": json.dumps(distractor_dict),
            "model_answer": "Option B is correct. list() creates generic vectors that accommodate heterogeneous types.",
            "marking_scheme": "2 marks for option B with valid identification.",
            "status": "approved",
            "is_validated": True,
            "source_filename": "Unit_1_Intro.pptx",
            "source_page_or_slide": "4",
            "source_topic": "Vectors and Lists"
        }
    ]

    # Test 1: Question Paper Mode (.docx)
    doc_path_qp = generate_question_paper_docx(sample_paper, questions, export_type="question_paper")
    doc_qp = docx.Document(doc_path_qp)
    full_text_qp = "\n".join(p.text for p in doc_qp.paragraphs)
    # Also include text inside table cells
    for table in doc_qp.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text_qp += "\n" + cell.text
    
    assert "R Programming for Data Analytics" in full_text_qp
    assert ("A. vector()" in full_text_qp) or ("A) vector()" in full_text_qp)
    assert ("B. list()" in full_text_qp) or ("B) list()" in full_text_qp)
    assert ("C. matrix()" in full_text_qp) or ("C) matrix()" in full_text_qp)
    assert ("D. array()" in full_text_qp) or ("D) array()" in full_text_qp)
    # In student question paper mode, correct answer or distractor rationales must NOT be displayed
    assert "list() creates generic vectors holding elements" not in full_text_qp

    # Test 2: Answer Key & Rubric Mode (.docx)
    doc_path_ak = generate_question_paper_docx(sample_paper, questions, export_type="with_answer_key")
    doc_ak = docx.Document(doc_path_ak)
    full_text_ak = "\n".join(p.text for p in doc_ak.paragraphs)
    for table in doc_ak.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text_ak += "\n" + cell.text

    assert ("Correct Answer: B" in full_text_ak) or ("Correct Option: B" in full_text_ak)
    assert "Distractor Rationales:" in full_text_ak
    assert "Option A:" in full_text_ak
    assert "Option C:" in full_text_ak
