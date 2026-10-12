import json
import pytest
import docx
from app.schemas.question_paper import QuestionResponse, McqOptionSchema
from app.services.semantic_validator import (
    validate_mcq_structure,
    evaluate_semantic_quality,
    validate_mcq_quality
)
from app.services.word_template_service import generate_question_paper_docx
from app.services.question_generator import R_PROGRAMMING_QUESTION_BANK

def test_mcq_canonical_structure_validity():
    """Verify that validate_mcq_structure accepts valid canonical MCQs and rejects invalid ones."""
    # 1. Valid Canonical MCQ
    valid_q = {
        "question_text": "Which operator is used to perform standard matrix multiplication in R?",
        "question_type": "mcq_single",
        "options_json": json.dumps([
            {"label": "A", "text": "*", "is_correct": False},
            {"label": "B", "text": "%*%", "is_correct": True},
            {"label": "C", "text": "%.%", "is_correct": False},
            {"label": "D", "text": "**", "is_correct": False}
        ]),
        "correct_answer_json": "B",
        "explanation": "%*% represents standard matrix inner product multiplication in base R."
    }
    valid, score, opt_grammar, issues = validate_mcq_structure(valid_q)
    assert valid is True
    assert score == 100.0
    assert len(issues) == 0

    # 2. Defect: Blank option text
    blank_opt_q = {
        "question_text": "Which operator is used to perform standard matrix multiplication in R?",
        "question_type": "mcq_single",
        "options_json": json.dumps([
            {"label": "A", "text": "*", "is_correct": False},
            {"label": "B", "text": "", "is_correct": True},
            {"label": "C", "text": "%.%", "is_correct": False},
            {"label": "D", "text": "**", "is_correct": False}
        ]),
        "correct_answer_json": "B",
        "explanation": "%*% represents standard matrix inner product multiplication in base R."
    }
    valid, score, _, issues = validate_mcq_structure(blank_opt_q)
    assert valid is False
    assert score == 0.0
    assert any("blank" in iss.lower() for iss in issues)

    # 3. Defect: Incorrect option count (e.g. 3 options)
    three_opts_q = {
        "question_text": "Which operator is used to perform standard matrix multiplication in R?",
        "question_type": "mcq_single",
        "options_json": json.dumps([
            {"label": "A", "text": "*", "is_correct": False},
            {"label": "B", "text": "%*%", "is_correct": True},
            {"label": "C", "text": "%.%", "is_correct": False}
        ]),
        "correct_answer_json": "B",
        "explanation": "Valid explanation."
    }
    valid, score, _, issues = validate_mcq_structure(three_opts_q)
    assert valid is False
    assert any("count" in iss.lower() for iss in issues)

    # 4. Defect: Multiple correct options in single-answer MCQ
    multi_correct_q = {
        "question_text": "Which operator is used to perform standard matrix multiplication in R?",
        "question_type": "mcq_single",
        "options_json": json.dumps([
            {"label": "A", "text": "*", "is_correct": True},
            {"label": "B", "text": "%*%", "is_correct": True},
            {"label": "C", "text": "%.%", "is_correct": False},
            {"label": "D", "text": "**", "is_correct": False}
        ]),
        "correct_answer_json": "B",
        "explanation": "Valid explanation."
    }
    valid, score, _, issues = validate_mcq_structure(multi_correct_q)
    assert valid is False
    assert any("more than one correct option" in iss.lower() for iss in issues)

    # 5. Defect: Zero correct options
    zero_correct_q = {
        "question_text": "Which operator is used to perform standard matrix multiplication in R?",
        "question_type": "mcq_single",
        "options_json": json.dumps([
            {"label": "A", "text": "*", "is_correct": False},
            {"label": "B", "text": "%*%", "is_correct": False},
            {"label": "C", "text": "%.%", "is_correct": False},
            {"label": "D", "text": "**", "is_correct": False}
        ]),
        "correct_answer_json": "",
        "explanation": "Valid explanation."
    }
    valid, score, _, issues = validate_mcq_structure(zero_correct_q)
    assert valid is False
    assert any("missing correct answer" in iss.lower() for iss in issues)

    # 6. Defect: Duplicate options
    duplicate_opts_q = {
        "question_text": "Which operator is used to perform standard matrix multiplication in R?",
        "question_type": "mcq_single",
        "options_json": json.dumps([
            {"label": "A", "text": "%*%", "is_correct": True},
            {"label": "B", "text": "%*%", "is_correct": False},
            {"label": "C", "text": "%.%", "is_correct": False},
            {"label": "D", "text": "**", "is_correct": False}
        ]),
        "correct_answer_json": "A",
        "explanation": "Valid explanation."
    }
    valid, score, _, issues = validate_mcq_structure(duplicate_opts_q)
    assert valid is False
    assert any("duplicate" in iss.lower() for iss in issues)


def test_mcq_semantic_validator_clamps_quality_to_zero():
    """Verify that any MCQ failing structural validation is forced to quality_score = 0.0 and status = needs_review."""
    invalid_q = {
        "question_text": "What is the result of adding two numeric vectors of unequal lengths in R?",
        "question_type": "mcq_single",
        "marks": 2,
        "co_level": "CO1",
        "bloom_level": "L2 - Understand",
        "difficulty": "Easy",
        "unit_number": "1",
        "source_topic": "Vectors and Data Types",
        "options_json": json.dumps([
            {"label": "A", "text": "", "is_correct": False},  # Empty option text!
            {"label": "B", "text": "Shorter vector recycled", "is_correct": True},
            {"label": "C", "text": "NA appended", "is_correct": False},
            {"label": "D", "text": "Truncated", "is_correct": False}
        ]),
        "correct_answer_json": "B",
        "explanation": "R recycles elements of the shorter vector.",
        "model_answer": "Option B is correct.",
        "marking_scheme": "2 marks for option B."
    }
    unit_info = {"number": "1", "title": "Unit 1", "topics": "Vectors and Data Types"}
    co_info = {"code": "CO1", "title": "Unit 1"}

    res = evaluate_semantic_quality(invalid_q, unit_info, co_info)
    assert res["quality_score"] == 0.0
    assert res["status"] == "needs_review"
    assert any("One or more MCQ options are blank" in r for r in res["reasons"])
    
    details = json.loads(res["details_json"])
    assert details["structural_validity"] is False
    assert details["overall_quality"] == 0.0


def test_mcq_api_serialization():
    """Verify that QuestionResponse serializes options_json into canonical list of McqOptionSchema with boolean is_correct."""
    raw_data = {
        "id": 101,
        "question_paper_id": 1,
        "section_name": "Part A",
        "question_number": "1",
        "question_text": "What is the primary role of the `levels` attribute in an R factor?",
        "question_type": "mcq_single",
        "marks": 2,
        "co_level": "CO2",
        "bloom_level": "L2 - Understand",
        "difficulty": "Easy",
        "status": "draft",
        "is_validated": True,
        "options_json": json.dumps([
            {"label": "A", "text": "Memory pointers", "is_correct": False},
            {"label": "B", "text": "Defines valid discrete categorical values", "is_correct": True},
            {"label": "C", "text": "Floating point operations", "is_correct": False},
            {"label": "D", "text": "Probability densities", "is_correct": False}
        ]),
        "correct_answer_json": "B",
        "explanation": "Levels define the allowable category values."
    }

    resp = QuestionResponse.model_validate(raw_data)
    assert resp.options is not None
    assert len(resp.options) == 4
    assert resp.correct_answer == "B"

    # Verify each option is an McqOptionSchema with strict boolean is_correct
    for opt in resp.options:
        assert isinstance(opt, McqOptionSchema)
        assert isinstance(opt.is_correct, bool)
        assert isinstance(opt.label, str)
        assert isinstance(opt.text, str)
        if opt.label == "B":
            assert opt.is_correct is True
        else:
            assert opt.is_correct is False


def test_five_sample_mcqs_from_specification():
    """
    Verify all 5 sample MCQs explicitly cited in the specification:
    1. Unequal vector addition / vector recycling
    2. Difference between NA and NULL
    3. Apply summary function column mean across matrix
    4. Matrix multiplication operator %*%
    5. Role of levels attribute in factor variable
    """
    # Look up in bank or define canonical representations
    samples = [
        # Sample 1
        {
            "topic": "Vectors",
            "stem": "In R, what is the result of adding two numeric vectors of unequal lengths, `x <- c(1, 2)` and `y <- c(10, 20, 30, 40)`?",
            "options": [
                {"label": "A", "text": "An error is raised due to incompatible vector dimensions.", "is_correct": False},
                {"label": "B", "text": "`c(11, 22, 31, 42)` because the shorter vector is recycled.", "is_correct": True},
                {"label": "C", "text": "`c(11, 22, 10, 20)` because NA values are appended to the shorter vector.", "is_correct": False},
                {"label": "D", "text": "`c(11, 22)` as only elements up to the length of the shorter vector are added.", "is_correct": False}
            ],
            "correct": "B",
            "explanation": "R automatically recycles elements of the shorter vector (c(1, 2, 1, 2)) to match the longer vector."
        },
        # Sample 2
        {
            "topic": "Missing Values",
            "stem": "What is the key functional difference between `NA` and `NULL` in R data structures?",
            "options": [
                {"label": "A", "text": "`NA` represents a missing or unknown value of a specific atomic type, whereas `NULL` represents the complete absence of an object or empty vector.", "is_correct": True},
                {"label": "B", "text": "`NA` can only be stored in character vectors, while `NULL` is reserved for numerical lists.", "is_correct": False},
                {"label": "C", "text": "`NA` terminates loop execution immediately, whereas `NULL` skips the current iteration.", "is_correct": False},
                {"label": "D", "text": "`NA` occupies 64 bytes of memory, while `NULL` dynamically allocates 128 bytes.", "is_correct": False}
            ],
            "correct": "A",
            "explanation": "NA represents missingness with an atomic type indicator, whereas NULL represents the empty object."
        },
        # Sample 3
        {
            "topic": "Apply Functions",
            "stem": "To apply a summary function `mean` across every column of a numeric matrix `X`, which `apply()` call is syntactically and semantically correct?",
            "options": [
                {"label": "A", "text": "`apply(X, MARGIN = 1, FUN = mean)`", "is_correct": False},
                {"label": "B", "text": "`apply(X, MARGIN = 2, FUN = mean)`", "is_correct": True},
                {"label": "C", "text": "`apply(X, MARGIN = c(1, 2), FUN = mean)`", "is_correct": False},
                {"label": "D", "text": "`lapply(X, FUN = mean)`", "is_correct": False}
            ],
            "correct": "B",
            "explanation": "In apply(), MARGIN = 1 indicates rows and MARGIN = 2 indicates columns."
        },
        # Sample 4
        {
            "topic": "Matrices",
            "stem": "Which operator is used to perform standard matrix multiplication between two compatible numeric matrices in R?",
            "options": [
                {"label": "A", "text": "`*`", "is_correct": False},
                {"label": "B", "text": "`%*%`", "is_correct": True},
                {"label": "C", "text": "`%.%`", "is_correct": False},
                {"label": "D", "text": "`**`", "is_correct": False}
            ],
            "correct": "B",
            "explanation": "%*% is the matrix product operator, while * performs element-wise multiplication."
        },
        # Sample 5
        {
            "topic": "Factors",
            "stem": "What is the primary role of the `levels` attribute in an R factor variable?",
            "options": [
                {"label": "A", "text": "It stores the underlying memory addresses of the data frame.", "is_correct": False},
                {"label": "B", "text": "It defines the complete set of valid discrete categorical values that the factor can represent.", "is_correct": True},
                {"label": "C", "text": "It determines whether the factor allows floating-point numerical operations.", "is_correct": False},
                {"label": "D", "text": "It converts the categorical values into continuous probability densities.", "is_correct": False}
            ],
            "correct": "B",
            "explanation": "The levels attribute defines the allowed categories mapped to underlying integer storage."
        }
    ]

    for sample in samples:
        # Check stem wording does not have banned enterprise fluff
        assert "enterprise" not in sample["stem"].lower()
        assert "scalability" not in sample["stem"].lower()
        assert "mission-critical" not in sample["stem"].lower()

        # Check structure
        q_dict = {
            "question_text": sample["stem"],
            "question_type": "mcq_single",
            "options_json": json.dumps(sample["options"]),
            "correct_answer_json": sample["correct"],
            "explanation": sample["explanation"]
        }
        valid, score, opt_grammar, issues = validate_mcq_structure(q_dict)
        assert valid is True, f"Sample {sample['topic']} failed validity: {issues}"
        assert score == 100.0
        assert opt_grammar >= 90.0

        # Check exactly one correct option with boolean True
        corrects = [opt for opt in sample["options"] if opt["is_correct"] is True]
        assert len(corrects) == 1
        assert corrects[0]["label"] == sample["correct"]


def test_mcq_word_export_formatting_and_student_privacy():
    """
    Verify Word document export:
    1. Question Paper mode (.docx):
       - Renders options A. - D. cleanly indented beneath the question stem.
       - Contains w:cantSplit on table rows to avoid page splits.
       - Does NOT expose correct answers, badges, or explanations to students.
    2. Answer Key mode (.docx):
       - Displays Correct Answer: [key]
       - Displays Explanation: [text]
    """
    sample_paper = {
        "id": 99,
        "title": "Semester Examination",
        "institution_name": "Department of Computer Science",
        "course_name": "R Programming for Data Analytics",
        "course_code": "CS302",
        "academic_year": "2025-2026",
        "semester": "IV",
        "duration": "3 Hours",
        "max_marks": 100,
        "instructions": "Answer all questions.",
        "programme": "B.Tech Computer Science",
        "regulation": "R2024",
        "generation_mode": "syllabus_ppt"
    }

    sample_questions = [
        {
            "id": 1,
            "section_name": "Part A",
            "question_number": "1",
            "question_text": "Which operator is used to perform standard matrix multiplication between two compatible numeric matrices in R?",
            "question_type": "mcq_single",
            "marks": 2,
            "unit_number": 2,
            "co_level": "CO2",
            "bloom_level": "L2 - Understand",
            "difficulty": "Easy",
            "options_json": json.dumps([
                {"label": "A", "text": "`*`", "is_correct": False},
                {"label": "B", "text": "`%*%`", "is_correct": True},
                {"label": "C", "text": "`%.%`", "is_correct": False},
                {"label": "D", "text": "`**`", "is_correct": False}
            ]),
            "correct_answer_json": "B",
            "correct_option": "B",
            "explanation": "%*% is the standard matrix multiplication operator in R.",
            "model_answer": "Correct Option: B.\n\nExplanation: %*% is the standard matrix multiplication operator in R.",
            "marking_scheme": "2 Marks: Selecting the correct option (B).",
            "status": "approved",
            "is_validated": True,
            "source_filename": "Unit_2_Matrices.pptx",
            "source_page_or_slide": "12",
            "source_topic": "Matrices and Arrays"
        }
    ]

    # 1. Question Paper Mode (Student View)
    doc_path_qp = generate_question_paper_docx(sample_paper, sample_questions, export_type="question_paper")
    doc_qp = docx.Document(doc_path_qp)

    # Verify cantSplit is present on question row
    found_cant_split = False
    for table in doc_qp.tables:
        for row in table.rows:
            trPr = row._tr.get_or_add_trPr()
            if trPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cantSplit") is not None:
                found_cant_split = True
                break
    assert found_cant_split is True, "Expected <w:cantSplit/> in table rows to prevent page break splitting"

    # Extract all text from paragraphs and table cells
    qp_text = "\n".join(p.text for p in doc_qp.paragraphs)
    for table in doc_qp.tables:
        for row in table.rows:
            for cell in row.cells:
                qp_text += "\n" + cell.text

    # Options must be present cleanly
    assert "A. `*`" in qp_text or "A) `*`" in qp_text
    assert "B. `%*%`" in qp_text or "B) `%*%`" in qp_text
    assert "C. `%.%`" in qp_text or "C) `%.%`" in qp_text
    assert "D. `**`" in qp_text or "D) `**`" in qp_text

    # Student View must NOT expose correct answers or explanations
    assert "Correct Answer: B" not in qp_text
    assert "Explanation: %*%" not in qp_text
    assert "Correct Option: B" not in qp_text
    assert "CORRECT" not in qp_text

    # 2. Answer Key Mode
    doc_path_ak = generate_question_paper_docx(sample_paper, sample_questions, export_type="with_answer_key")
    doc_ak = docx.Document(doc_path_ak)
    ak_text = "\n".join(p.text for p in doc_ak.paragraphs)
    for table in doc_ak.tables:
        for row in table.rows:
            for cell in row.cells:
                ak_text += "\n" + cell.text

    assert "Correct Answer: B" in ak_text
    assert "Explanation:" in ak_text
    assert "%*% is the standard matrix multiplication operator in R" in ak_text
