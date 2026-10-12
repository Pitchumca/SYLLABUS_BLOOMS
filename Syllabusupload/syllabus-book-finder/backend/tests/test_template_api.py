import pytest
from pathlib import Path
from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)

@pytest.fixture
def auth_header():
    # Login as default faculty
    res = client.post("/api/auth/login", json={"email": "faculty@college.edu", "password": "Faculty@123"})
    assert res.status_code == 200, res.text
    token = res.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}

def test_template_api_workflow(auth_header, tmp_path: Path):
    # 1. Create a sample docx template
    import docx
    doc_path = tmp_path / "ubit_sample.docx"
    doc = docx.Document()
    doc.add_paragraph("AARUPADAI VEEDU INSTITUTE OF TECHNOLOGY")
    doc.add_paragraph("Course Code: 35021P28 | Course Title: R Programming")
    doc.add_paragraph("Max. Marks: 50 | Duration: 90 Minutes")
    
    doc.add_paragraph("PART - A")
    tbl_a = doc.add_table(rows=1, cols=5)
    hdr_a = tbl_a.rows[0].cells
    hdr_a[0].text = "Q.No"
    hdr_a[1].text = "Questions"
    hdr_a[2].text = "CO"
    hdr_a[3].text = "Bloom's Level"
    hdr_a[4].text = "Marks"

    for i in range(1, 6):
        r = tbl_a.add_row().cells
        r[0].text = str(i)
        r[1].text = f"Explain concept {i}"
        r[2].text = "CO1"
        r[3].text = "L2"
        r[4].text = "2"

    doc.add_paragraph("PART - B")
    tbl_b = doc.add_table(rows=1, cols=5)
    hdr_b = tbl_b.rows[0].cells
    hdr_b[0].text = "Q.No"
    hdr_b[1].text = "Questions"
    hdr_b[2].text = "CO"
    hdr_b[3].text = "Bloom's Level"
    hdr_b[4].text = "Marks"

    # 4 questions of 10 marks
    for i in range(6, 10):
        r = tbl_b.add_row().cells
        r[0].text = str(i)
        r[1].text = f"Solve programming problem {i}"
        r[2].text = "CO2"
        r[3].text = "L3"
        r[4].text = "10"

    doc.save(str(doc_path))

    # 2. Upload template to /api/template/upload
    with open(doc_path, "rb") as f:
        res_upload = client.post(
            "/api/template/upload",
            files={"file": ("ubit_sample.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            headers=auth_header
        )
    assert res_upload.status_code == 200
    upload_data = res_upload.json()
    assert "blueprint" in upload_data
    assert "confidence_scores" in upload_data
    assert upload_data["blueprint"]["maximum_marks"] == 50
    assert len(upload_data["blueprint"]["sections"]) == 2

    # 3. Validate Blueprint endpoint
    bp = upload_data["blueprint"]
    res_val = client.post("/api/template/validate", json=bp, headers=auth_header)
    assert res_val.status_code == 200
    assert res_val.json()["is_valid"] is True
    assert res_val.json()["calculated_total_marks"] == 50

    # 4. Save to Template Library
    save_payload = {
        "template_name": "AVIT 50 Marks Test Template",
        "institution": bp["institution"],
        "course_code": bp["course_code"],
        "course_title": bp["course_title"],
        "maximum_marks": bp["maximum_marks"],
        "duration_minutes": bp["duration_minutes"],
        "blueprint": bp,
        "original_filename": "ubit_sample.docx",
        "stored_filename": upload_data["stored_filename"],
        "file_format": "docx"
    }
    res_save = client.post("/api/template/save", json=save_payload, headers=auth_header)
    assert res_save.status_code == 201
    saved_template = res_save.json()
    template_id = saved_template["id"]
    assert saved_template["name"] == "AVIT 50 Marks Test Template"
    assert saved_template["version"] == 1

    # 5. List Library
    res_lib = client.get("/api/template/library", headers=auth_header)
    assert res_lib.status_code == 200
    assert any(t["id"] == template_id for t in res_lib.json())

    # 6. Retrieve Template Detail
    res_detail = client.get(f"/api/template/{template_id}", headers=auth_header)
    assert res_detail.status_code == 200
    assert res_detail.json()["id"] == template_id
    assert res_detail.json()["blueprint"]["maximum_marks"] == 50

    # 7. Update Template (Increment Version)
    update_payload = {
        "template_name": "AVIT 50 Marks Test Template (Updated)",
        "change_summary": "Revised title and instructions"
    }
    res_update = client.put(f"/api/template/{template_id}", json=update_payload, headers=auth_header)
    assert res_update.status_code == 200
    assert res_update.json()["version"] == 2
    assert res_update.json()["name"] == "AVIT 50 Marks Test Template (Updated)"

    # 8. Export Blueprint JSON
    res_export = client.get(f"/api/template/{template_id}/export-blueprint-json", headers=auth_header)
    assert res_export.status_code == 200
    assert res_export.headers["content-type"] == "application/json"
    exported_data = res_export.json()
    assert exported_data["maximum_marks"] == 50
